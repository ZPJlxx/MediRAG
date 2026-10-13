"""文档加载与切分（迁移自 general-doc-qa-assistant 的 custom_loader.py）。

保留其核心价值：
- 按文件类型选择合适的 Loader（.docx 本地解析避免 unstructured 重依赖）
- 中文感知的递归字符切分
- location_label 证据位置标签（页码/行号/章节路径）——医疗引用溯源的基础

调整：.doc 老格式给出明确提示（建议转存 .docx），不引入 unstructured。
"""

import bisect
import os
import re
from enum import Enum
from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from loguru import logger

from app.core.config import settings


class FileType(Enum):
    CSV = "csv"
    TXT = "txt"
    DOCX = "docx"
    PDF = "pdf"
    MD = "md"


_SUFFIX_TO_FILETYPE = {
    "csv": FileType.CSV,
    "txt": FileType.TXT,
    "docx": FileType.DOCX,
    "pdf": FileType.PDF,
    "md": FileType.MD,
}

SUPPORTED_KNOWLEDGE_SUFFIXES = {f".{t.value}" for t in FileType}


def detect_filetype(file_path):
    suffix = Path(file_path).suffix.lower().lstrip(".")
    return _SUFFIX_TO_FILETYPE.get(suffix)


class LocalDocxLoader(BaseLoader):
    """本地解析 .docx（python-docx），提取正文/表格 + 章节路径元数据。"""

    def __init__(self, file_path: str):
        self.file_path = file_path

    @staticmethod
    def _resolve_heading_level(paragraph, text):
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        heading_match = re.match(r"Heading\s*(\d+)", style_name, re.IGNORECASE)
        if heading_match:
            return max(1, int(heading_match.group(1)))
        if re.match(r"^[一二三四五六七八九十]+、", text) and (len(text) <= 20 or text.endswith("：")):
            return 1
        if re.match(r"^\d+、", text) and (len(text) <= 28 or text.endswith("：")):
            return 2
        numbered_match = re.match(r"^(\d+(?:\.\d+)+)", text)
        if numbered_match and (len(text) <= 18 or text.endswith("：")):
            return min(6, numbered_match.group(1).count(".") + 2)
        return None

    @staticmethod
    def _extract_inline_label(text):
        if "：" not in text:
            return None
        prefix = text.split("：", 1)[0].strip()
        if 2 <= len(prefix) <= 20 and len(prefix) <= len(text) * 0.4:
            return prefix
        return None

    def load(self):
        doc = DocxDocument(self.file_path)
        documents = []
        section_stack = []
        section_block_index = 0
        block_index = 0

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            heading_level = self._resolve_heading_level(paragraph, text)
            anchor_label = self._extract_inline_label(text)

            if heading_level is not None:
                section_stack = section_stack[: heading_level - 1]
                section_stack.append(text)
                section_block_index = 0
            else:
                section_block_index += 1

            block_index += 1
            metadata = {
                "source": str(self.file_path),
                "block_index": block_index,
                "section_path": " > ".join(section_stack),
            }
            if heading_level is not None:
                metadata["anchor_label"] = text
                metadata["section_block_index"] = 0
            else:
                metadata["section_block_index"] = section_block_index
                if anchor_label:
                    metadata["anchor_label"] = anchor_label

            documents.append(Document(page_content=text, metadata=metadata))

        for table_index, table in enumerate(doc.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if not row_text:
                    continue
                section_block_index += 1
                block_index += 1
                documents.append(
                    Document(
                        page_content=row_text,
                        metadata={
                            "source": str(self.file_path),
                            "block_index": block_index,
                            "section_path": " > ".join(section_stack),
                            "section_block_index": section_block_index,
                            "anchor_label": f"表格{table_index} 第{row_index}行",
                        },
                    )
                )

        if documents:
            return documents
        return [Document(page_content="", metadata={"source": str(self.file_path)})]


class MyCustomLoader(BaseLoader):
    """根据文件类型选择加载器，并统一做中文感知切分与位置标签富化。"""

    file_type = {
        FileType.CSV: (CSVLoader, {"autodetect_encoding": True}),
        FileType.TXT: (TextLoader, {"encoding": "utf-8"}),
        FileType.DOCX: (LocalDocxLoader, {}),
        FileType.PDF: (PyPDFLoader, {}),
        FileType.MD: (TextLoader, {"encoding": "utf-8"}),
    }

    def __init__(self, file_path: str, chunk_size: int | None = None, chunk_overlap: int | None = None):
        self.file_path = Path(file_path)
        self.detected_type = detect_filetype(file_path)
        if self.detected_type is None:
            suffix = self.file_path.suffix.lower()
            if suffix == ".doc":
                raise ValueError(
                    f"暂不支持 Word 97-2003 格式（{self.file_path.name}），请另存为 .docx 后重新上传"
                )
            raise ValueError(f"无法识别的文档类型：{suffix or '(无后缀)'}")
        loader_class, params = self.file_type[self.detected_type]
        logger.info(f"知识文档[{self.file_path.name}]使用加载器: {loader_class.__name__}")
        self.loader: BaseLoader = loader_class(file_path, **params)

        self.text_splitter = RecursiveCharacterTextSplitter(
            separators=["\n\n", "\n", "。|！|？", r"\.\s|\!\s|\?\s", "；|;\\s", "，|,\\s", " ", ""],
            chunk_size=chunk_size or settings.chunk_size,
            chunk_overlap=chunk_overlap or settings.chunk_overlap,
            length_function=len,
            add_start_index=True,
        )
        self._raw_file_text = None
        self._line_offsets = None

    def _get_raw_file_text(self):
        if self._raw_file_text is not None:
            return self._raw_file_text
        self._raw_file_text = self.file_path.read_text(encoding="utf-8", errors="ignore")
        return self._raw_file_text

    def _get_line_offsets(self):
        if self._line_offsets is not None:
            return self._line_offsets
        raw_text = self._get_raw_file_text()
        offsets = [0]
        position = 0
        for line in raw_text.splitlines(keepends=True):
            position += len(line)
            offsets.append(position)
        self._line_offsets = offsets
        return offsets

    def _build_line_range_label(self, document):
        start_index = document.metadata.get("start_index")
        chunk_text = getattr(document, "page_content", "") or ""
        raw_text = self._get_raw_file_text()

        if start_index is None and chunk_text:
            start_index = raw_text.find(chunk_text)
        if start_index is None or start_index < 0:
            return None

        end_index = start_index + max(1, len(chunk_text))
        line_offsets = self._get_line_offsets()
        start_line = bisect.bisect_right(line_offsets, start_index)
        end_line = bisect.bisect_right(line_offsets, max(start_index, end_index - 1))

        line_label = f"行 {start_line}" if start_line == end_line else f"行 {start_line}-{end_line}"
        if self.detected_type != FileType.MD:
            return line_label

        markdown_lines = raw_text.splitlines()
        heading_label = None
        for line_number in range(min(start_line - 1, len(markdown_lines) - 1), -1, -1):
            stripped_line = markdown_lines[line_number].strip()
            if stripped_line.startswith("#"):
                heading_label = stripped_line.lstrip("#").strip()
                break
        if heading_label:
            return f"{heading_label} · {line_label}"
        return line_label

    @staticmethod
    def _build_docx_location_label(document):
        metadata = getattr(document, "metadata", {}) or {}
        section_path = metadata.get("section_path", "").strip()
        anchor_label = metadata.get("anchor_label", "").strip()
        section_block_index = metadata.get("section_block_index")

        if section_path:
            section_parts = [part.strip() for part in section_path.split(" > ") if part.strip()]
            section_path = " > ".join(section_parts[-3:])

        if anchor_label and section_path:
            last_section = section_path.split(" > ")[-1]
            if anchor_label != last_section:
                return f"{section_path} · {anchor_label}"
            return section_path
        if section_path and section_block_index:
            return f"{section_path} · 第{section_block_index}段附近"
        if anchor_label:
            return anchor_label
        if section_path:
            return section_path
        return None

    def _build_location_label(self, document):
        metadata = getattr(document, "metadata", {}) or {}

        if self.detected_type == FileType.PDF:
            page_number = metadata.get("page", metadata.get("page_number"))
            if page_number is not None:
                return f"第 {int(page_number) + 1} 页"
            return None

        if self.detected_type == FileType.CSV:
            row_number = metadata.get("row", metadata.get("row_index"))
            if row_number is not None:
                return f"第 {int(row_number) + 1} 行"
            return None

        if self.detected_type in {FileType.TXT, FileType.MD}:
            return self._build_line_range_label(document)

        if self.detected_type == FileType.DOCX:
            return self._build_docx_location_label(document)

        return None

    def _enrich_documents(self, documents):
        enriched_documents = []
        for document in documents:
            enriched_document = Document(
                page_content=document.page_content,
                metadata=dict(getattr(document, "metadata", {}) or {}),
            )
            location_label = self._build_location_label(enriched_document)
            if location_label:
                enriched_document.metadata["location_label"] = location_label
            enriched_documents.append(enriched_document)
        return enriched_documents

    def lazy_load(self):
        raw_documents = self.loader.load()
        documents = self.text_splitter.split_documents(raw_documents)
        return self._enrich_documents(documents)

    def load(self):
        return self.lazy_load()
